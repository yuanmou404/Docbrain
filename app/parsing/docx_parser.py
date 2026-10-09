"""Word(.docx) 解析：段落样式识别标题 + 表格提取。"""
import logging

import docx

from app.parsing.chunker import Block

logger = logging.getLogger("docbrain.docx")


def parse_docx(path: str) -> tuple[list[Block], dict]:
    d = docx.Document(path)
    blocks: list[Block] = []

    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style_name = (p.style.name if p.style is not None else "") or ""
        is_heading = style_name.startswith("Heading") or "标题" in style_name
        blocks.append(Block(text=text, heading=text if is_heading else None))

    # 表格追加为独立块（Word 无页码概念，page=None）
    for tbl in d.tables:
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in tbl.rows]
        rows = [r for r in rows if r.strip(" |")]
        if rows:
            blocks.append(Block(text="\n".join(rows), heading=None))

    return blocks, {"pages": None, "ocr_pages": 0}
